#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
实验结果导出为Excel表格
自动解析JSON结果文件，生成多Sheet的Excel报告
"""

import json
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
import numpy as np
from datetime import datetime


def load_result_json(json_path):
    """加载JSON结果文件"""
    with open(json_path, 'r', encoding='utf-8') as f:
        return json.load(f)


def create_accuracy_comparison_sheet(wb, results):
    """创建准确率对比表"""
    ws = wb.create_sheet("准确率对比", 0)
    
    # 标题
    ws['A1'] = '防御方法对比实验 - 准确率'
    ws['A1'].font = Font(size=14, bold=True)
    ws.merge_cells('A1:H1')
    
    # 表头
    headers = ['轮次', 'TEE-FL准确率', 'TEE-FL损失', 'FLTrust准确率', 'FLTrust损失', 
               'FedAvg准确率', 'FedAvg损失', '最佳方法']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(3, col, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal='center')
    
    # 数据
    max_epochs = max(len(results['tee_fl']['accuracies']), 
                     len(results['fltrust']['accuracies']),
                     len(results['fedavg']['accuracies']))
    
    for i in range(max_epochs):
        row = 4 + i
        ws.cell(row, 1, i+1)  # 轮次
        
        # TEE-FL
        if i < len(results['tee_fl']['accuracies']):
            tee_acc = results['tee_fl']['accuracies'][i]['accuracy']
            tee_loss = results['tee_fl']['accuracies'][i]['loss']
            ws.cell(row, 2, f"{tee_acc:.2f}%")
            ws.cell(row, 3, f"{tee_loss:.4f}")
        
        # FLTrust
        if i < len(results['fltrust']['accuracies']):
            flt_acc = results['fltrust']['accuracies'][i]['accuracy']
            flt_loss = results['fltrust']['accuracies'][i]['loss']
            ws.cell(row, 4, f"{flt_acc:.2f}%")
            ws.cell(row, 5, f"{flt_loss:.4f}")
        
        # FedAvg
        if i < len(results['fedavg']['accuracies']):
            fed_acc = results['fedavg']['accuracies'][i]['accuracy']
            fed_loss = results['fedavg']['accuracies'][i]['loss']
            ws.cell(row, 6, f"{fed_acc:.2f}%")
            ws.cell(row, 7, f"{fed_loss:.4f}")
        
        # 最佳方法
        accs = []
        if i < len(results['tee_fl']['accuracies']):
            accs.append(('TEE-FL', tee_acc))
        if i < len(results['fltrust']['accuracies']):
            accs.append(('FLTrust', flt_acc))
        if i < len(results['fedavg']['accuracies']):
            accs.append(('FedAvg', fed_acc))
        
        if accs:
            best = max(accs, key=lambda x: x[1])
            cell = ws.cell(row, 8, best[0])
            if best[0] == 'TEE-FL':
                cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    
    # 调整列宽
    ws.column_dimensions['A'].width = 8
    for col in ['B', 'C', 'D', 'E', 'F', 'G']:
        ws.column_dimensions[col].width = 14
    ws.column_dimensions['H'].width = 12


def create_summary_sheet(wb, results):
    """创建结果汇总表"""
    ws = wb.create_sheet("结果汇总", 1)
    
    # 标题
    ws['A1'] = '实验结果汇总'
    ws['A1'].font = Font(size=14, bold=True)
    ws.merge_cells('A1:E1')
    
    # 基本信息
    ws['A3'] = '实验配置'
    ws['A3'].font = Font(bold=True, size=12)
    ws.merge_cells('A3:B3')
    
    config = results.get('config', {})
    config_items = [
        ('数据集', config.get('dataset', 'N/A')),
        ('模型', config.get('model', 'N/A')),
        ('客户端数量', config.get('num_users', 'N/A')),
        ('训练轮数', config.get('epochs', 'N/A')),
        ('每轮采样率', f"{config.get('frac', 0)*100:.0f}%"),
        ('本地训练轮数', config.get('local_ep', 'N/A')),
        ('攻击类型', config.get('attack_type', 'N/A')),
        ('恶意比例', f"{config.get('malicious_ratio', 0)*100:.0f}%"),
        ('数据分布', 'Non-IID' if config.get('iid', 0) == 0 else 'IID'),
        ('Beta参数', config.get('data_beta', 'N/A')),
    ]
    
    row = 4
    for key, value in config_items:
        ws.cell(row, 1, key)
        ws.cell(row, 2, value)
        ws.cell(row, 1).font = Font(bold=True)
        row += 1
    
    # 性能对比
    ws.cell(row + 1, 1, '性能对比').font = Font(bold=True, size=12)
    ws.merge_cells(f'A{row+1}:E{row+1}')
    
    # 表头
    headers = ['方法', '最终准确率', '最高准确率', '最低准确率', '平均准确率']
    row += 2
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row, col, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal='center')
    
    # 计算统计数据
    methods_data = []
    for method_name, method_key in [('TEE-FL', 'tee_fl'), ('FLTrust', 'fltrust'), ('FedAvg', 'fedavg')]:
        if method_key in results:
            accs = [item['accuracy'] for item in results[method_key]['accuracies']]
            methods_data.append({
                '方法': method_name,
                '最终准确率': f"{accs[-1]:.2f}%",
                '最高准确率': f"{max(accs):.2f}%",
                '最低准确率': f"{min(accs):.2f}%",
                '平均准确率': f"{np.mean(accs):.2f}%"
            })
    
    row += 1
    for method_data in methods_data:
        for col, key in enumerate(headers, 1):
            cell = ws.cell(row, col, method_data.get(key, 'N/A'))
            if method_data['方法'] == 'TEE-FL':
                cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        row += 1
    
    # 训练时间
    row += 1
    ws.cell(row, 1, '训练时间').font = Font(bold=True, size=12)
    row += 1
    for method_name, method_key in [('TEE-FL', 'tee_fl'), ('FLTrust', 'fltrust'), ('FedAvg', 'fedavg')]:
        if method_key in results and 'times' in results[method_key]:
            time_sec = results[method_key]['times'][0]
            time_min = time_sec / 60
            time_hour = time_min / 60
            ws.cell(row, 1, method_name)
            ws.cell(row, 2, f"{time_hour:.2f}小时 ({time_min:.1f}分钟)")
            ws.cell(row, 1).font = Font(bold=True)
            row += 1
    
    # 调整列宽
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 18
    for col in ['C', 'D', 'E']:
        ws.column_dimensions[col].width = 15


def create_detection_stats_sheet(wb, results):
    """创建检测统计表（TEE-FL）"""
    if 'tee_fl' not in results or 'detection_stats' not in results['tee_fl']:
        return
    
    ws = wb.create_sheet("检测性能", 2)
    
    # 标题
    ws['A1'] = 'TEE-FL 检测性能统计'
    ws['A1'].font = Font(size=14, bold=True)
    ws.merge_cells('A1:D1')
    
    stats = results['tee_fl']['detection_stats'][0]
    
    # 检测结果
    ws['A3'] = '检测结果汇总'
    ws['A3'].font = Font(bold=True, size=12)
    ws.merge_cells('A3:D3')
    
    detection_items = [
        ('总客户端数', stats.get('total_clients', 0)),
        ('实际恶意客户端', stats.get('true_malicious', 0)),
        ('检测为恶意', stats.get('detected_malicious', 0)),
        ('', ''),
        ('真阳性 (TP)', stats.get('true_positives', 0)),
        ('假阳性 (FP)', stats.get('false_positives', 0)),
        ('假阴性 (FN)', stats.get('false_negatives', 0)),
        ('', ''),
        ('召回率 (Recall)', f"{stats.get('recall', 0)*100:.2f}%"),
        ('精确率 (Precision)', f"{stats.get('precision', 0)*100:.2f}%"),
    ]
    
    row = 4
    for key, value in detection_items:
        if key:
            ws.cell(row, 1, key)
            ws.cell(row, 2, value)
            ws.cell(row, 1).font = Font(bold=True)
            if '率' in key:
                ws.cell(row, 2).fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        row += 1
    
    # 混淆矩阵
    row += 1
    ws.cell(row, 1, '混淆矩阵').font = Font(bold=True, size=12)
    row += 1
    
    # 表头
    ws.cell(row, 2, '预测为良性')
    ws.cell(row, 3, '预测为恶意')
    ws.cell(row, 2).font = Font(bold=True)
    ws.cell(row, 3).font = Font(bold=True)
    
    row += 1
    ws.cell(row, 1, '实际良性')
    ws.cell(row, 1).font = Font(bold=True)
    tn = stats['total_clients'] - stats['true_malicious'] - stats['false_positives']
    ws.cell(row, 2, tn)  # TN
    ws.cell(row, 3, stats['false_positives'])  # FP
    
    row += 1
    ws.cell(row, 1, '实际恶意')
    ws.cell(row, 1).font = Font(bold=True)
    ws.cell(row, 2, stats['false_negatives'])  # FN
    ws.cell(row, 3, stats['true_positives'])  # TP
    ws.cell(row, 3).fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    
    # 调整列宽
    for col in ['A', 'B', 'C', 'D']:
        ws.column_dimensions[col].width = 18


def create_trust_scores_sheet(wb, results):
    """创建FLTrust信任分数表"""
    if 'fltrust' not in results or 'trust_scores' not in results['fltrust']:
        return
    
    ws = wb.create_sheet("FLTrust信任分数", 3)
    
    # 标题
    ws['A1'] = 'FLTrust 信任分数统计'
    ws['A1'].font = Font(size=14, bold=True)
    ws.merge_cells('A1:F1')
    
    # 表头
    headers = ['轮次', '平均信任分数', '最小值', '最大值', '标准差', '零信任分数客户端数']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(3, col, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal='center')
    
    # 数据
    trust_scores = results['fltrust']['trust_scores']
    for i, epoch_scores in enumerate(trust_scores):
        row = 4 + i
        ws.cell(row, 1, epoch_scores['epoch'])
        ws.cell(row, 2, f"{epoch_scores['mean']:.4f}")
        ws.cell(row, 3, f"{epoch_scores['min']:.4f}")
        ws.cell(row, 4, f"{epoch_scores['max']:.4f}")
        ws.cell(row, 5, f"{epoch_scores['std']:.4f}")
        
        # 计算零信任分数客户端数
        zero_count = sum(1 for s in epoch_scores['scores'] if s == 0)
        ws.cell(row, 6, f"{zero_count}/{len(epoch_scores['scores'])}")
        
        # 低信任分数警告
        if epoch_scores['mean'] < 0.2:
            for col in range(1, 7):
                ws.cell(row, col).fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    
    # 调整列宽
    for col in ['A', 'B', 'C', 'D', 'E', 'F']:
        ws.column_dimensions[col].width = 18


def create_detailed_accuracy_sheet(wb, results):
    """创建详细准确率数据表"""
    ws = wb.create_sheet("详细准确率数据", 4)
    
    # 标题
    ws['A1'] = '详细准确率和损失数据'
    ws['A1'].font = Font(size=14, bold=True)
    ws.merge_cells('A1:G1')
    
    # TEE-FL数据
    ws['A3'] = 'TEE-FL'
    ws['A3'].font = Font(bold=True, size=12)
    ws['A3'].fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    
    headers = ['轮次', '准确率(%)', '损失']
    for col, header in enumerate(headers, 1):
        ws.cell(4, col, header).font = Font(bold=True)
    
    row = 5
    for item in results['tee_fl']['accuracies']:
        ws.cell(row, 1, item['epoch'])
        ws.cell(row, 2, item['accuracy'])
        ws.cell(row, 3, item['loss'])
        row += 1
    
    # FLTrust数据
    row += 1
    ws.cell(row, 1, 'FLTrust').font = Font(bold=True, size=12)
    ws.cell(row, 1).fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
    
    row += 1
    for col, header in enumerate(headers, 1):
        ws.cell(row, col, header).font = Font(bold=True)
    
    row += 1
    for item in results['fltrust']['accuracies']:
        ws.cell(row, 1, item['epoch'])
        ws.cell(row, 2, item['accuracy'])
        ws.cell(row, 3, item['loss'])
        row += 1
    
    # FedAvg数据
    row += 1
    ws.cell(row, 1, 'FedAvg').font = Font(bold=True, size=12)
    ws.cell(row, 1).fill = PatternFill(start_color="F4B084", end_color="F4B084", fill_type="solid")
    
    row += 1
    for col, header in enumerate(headers, 1):
        ws.cell(row, col, header).font = Font(bold=True)
    
    row += 1
    for item in results['fedavg']['accuracies']:
        ws.cell(row, 1, item['epoch'])
        ws.cell(row, 2, item['accuracy'])
        ws.cell(row, 3, item['loss'])
        row += 1
    
    # 调整列宽
    for col in ['A', 'B', 'C']:
        ws.column_dimensions[col].width = 15


def export_to_excel(json_path, output_path=None):
    """主函数：导出JSON结果到Excel"""
    # 加载JSON
    print(f"📖 读取结果文件: {json_path}")
    results = load_result_json(json_path)
    
    # 生成输出文件名
    if output_path is None:
        base_name = os.path.splitext(json_path)[0]
        output_path = f"{base_name}.xlsx"
    
    # 创建Excel工作簿
    print("📊 创建Excel工作簿...")
    wb = Workbook()
    wb.remove(wb.active)  # 删除默认sheet
    
    # 创建各个sheet
    print("  ✓ 准确率对比表")
    create_accuracy_comparison_sheet(wb, results)
    
    print("  ✓ 结果汇总表")
    create_summary_sheet(wb, results)
    
    print("  ✓ 检测性能表")
    create_detection_stats_sheet(wb, results)
    
    print("  ✓ FLTrust信任分数表")
    create_trust_scores_sheet(wb, results)
    
    print("  ✓ 详细准确率数据表")
    create_detailed_accuracy_sheet(wb, results)
    
    # 保存
    wb.save(output_path)
    print(f"✅ Excel文件已保存: {output_path}")
    
    return output_path


def main():
    """主程序"""
    import glob
    
    # 查找最新的结果文件
    result_dir = "results/defense_comparison"
    json_files = glob.glob(f"{result_dir}/comparison_*.json")
    
    if not json_files:
        print("❌ 未找到结果文件")
        print(f"   请确保 {result_dir} 目录下有 comparison_*.json 文件")
        return
    
    # 按修改时间排序，选择最新的
    latest_file = max(json_files, key=os.path.getmtime)
    
    print("=" * 80)
    print("📊 实验结果导出工具")
    print("=" * 80)
    print(f"找到 {len(json_files)} 个结果文件")
    print(f"最新文件: {latest_file}")
    print()
    
    # 导出
    output_file = export_to_excel(latest_file)
    
    print()
    print("=" * 80)
    print("✅ 导出完成！")
    print(f"📁 Excel文件: {output_file}")
    print("=" * 80)
    
    # 列出所有生成的sheet
    print("\n📋 包含的数据表:")
    print("  1. 准确率对比 - 各方法每轮准确率和损失对比")
    print("  2. 结果汇总 - 实验配置和性能统计")
    print("  3. 检测性能 - TEE-FL的检测统计和混淆矩阵")
    print("  4. FLTrust信任分数 - 每轮信任分数统计")
    print("  5. 详细准确率数据 - 原始数据完整记录")


if __name__ == "__main__":
    main()
